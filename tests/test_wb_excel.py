"""Tests for WB Excel parser/writer (app/services/excel/wb_excel.py)."""
import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from app.services.excel.wb_excel import (
    WB_PRODUCTS_SHEET,
    WB_SYSTEM_COLUMNS,
    WbExcelReader,
    WbExcelWriter,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_wb_excel(tmp_path: Path, rows: list[dict], add_instruction: bool = False) -> Path:
    """Create a minimal WB-style Excel file with a 'Товары' sheet."""
    path = tmp_path / "wb_template.xlsx"
    df = pd.DataFrame(rows)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=WB_PRODUCTS_SHEET, index=False)
        if add_instruction:
            instr_df = pd.DataFrame([["Инструкция строка 1"], ["Инструкция строка 2"]])
            instr_df.to_excel(writer, sheet_name="Инструкция", index=False, header=False)
    return path


# ---------------------------------------------------------------------------
# 1. Reader extracts products from template
# ---------------------------------------------------------------------------

def test_reader_extracts_products_from_template(tmp_path):
    rows = [
        {
            "Артикул продавца": "SKU-001",
            "Артикул WB": "WB-111",
            "Наименование": "Тестовый товар",
            "Бренд": "TestBrand",
            "Описание": "Описание товара",
            "Материал": "Хлопок",
            "Страна производства": "Россия",
        }
    ]
    path = _make_wb_excel(tmp_path, rows)
    reader = WbExcelReader(path)
    products = reader.read_products()

    assert len(products) == 1
    p = products[0]
    assert p["sku"] == "SKU-001"
    assert p["name"] == "Тестовый товар"
    assert p["brand"] == "TestBrand"
    assert p["description"] == "Описание товара"
    # Characteristics contain non-system columns
    assert "Материал" in p["characteristics"]
    assert p["characteristics"]["Материал"] == "Хлопок"
    assert "Страна производства" in p["characteristics"]


# ---------------------------------------------------------------------------
# 2. Reader skips empty rows
# ---------------------------------------------------------------------------

def test_reader_skips_empty_rows(tmp_path):
    rows = [
        {"Артикул продавца": "SKU-001", "Наименование": "Товар 1", "Материал": "Хлопок"},
        {"Артикул продавца": None, "Наименование": None, "Материал": None},  # empty row
        {"Артикул продавца": "SKU-002", "Наименование": "Товар 2", "Материал": "Полиэстер"},
    ]
    path = _make_wb_excel(tmp_path, rows)
    reader = WbExcelReader(path)
    products = reader.read_products()

    skus = [p["sku"] for p in products]
    assert "SKU-001" in skus
    assert "SKU-002" in skus
    # Empty row should be skipped
    assert len(products) == 2


# ---------------------------------------------------------------------------
# 3. Reader separates system columns from characteristics
# ---------------------------------------------------------------------------

def test_reader_separates_system_columns_from_characteristics(tmp_path):
    rows = [
        {
            "Артикул продавца": "SKU-001",
            "Артикул WB": "WB-111",
            "Наименование": "Товар",
            "Категория продавца": "Одежда",
            "Бренд": "Brand",
            "Размер": "M",
            "Цвет": "Красный",
            "Состав": "100% хлопок",
            "Описание": "Описание",
            "Группа": "Мужская одежда",
            # Actual characteristics
            "Материал": "Хлопок",
            "Сезон": "Лето",
        }
    ]
    path = _make_wb_excel(tmp_path, rows)
    reader = WbExcelReader(path)
    products = reader.read_products()

    assert len(products) == 1
    characteristics = products[0]["characteristics"]
    # System columns must NOT appear in characteristics
    for sys_col in WB_SYSTEM_COLUMNS:
        assert sys_col not in characteristics, f"System column '{sys_col}' leaked into characteristics"
    # Non-system columns must appear
    assert "Материал" in characteristics
    assert "Сезон" in characteristics


# ---------------------------------------------------------------------------
# 4. Writer preserves original structure (Инструкция sheet copied)
# ---------------------------------------------------------------------------

def test_writer_preserves_original_structure(tmp_path):
    rows = [
        {"Артикул продавца": "SKU-001", "Наименование": "Товар", "Материал": ""}
    ]
    path = _make_wb_excel(tmp_path, rows, add_instruction=True)

    reader = WbExcelReader(path)
    products = reader.read_products()
    products[0]["ai_filled"] = {"Материал": "Шерсть"}

    output = tmp_path / "output.xlsx"
    writer = WbExcelWriter(path)
    result = writer.write_filled(output, products)

    assert result.exists()
    # Both sheets must exist in output
    wb = load_workbook(result)
    assert WB_PRODUCTS_SHEET in wb.sheetnames
    assert "Инструкция" in wb.sheetnames


# ---------------------------------------------------------------------------
# 5. Writer only fills empty cells (never overwrites existing values)
# ---------------------------------------------------------------------------

def test_writer_only_fills_empty_cells(tmp_path):
    rows = [
        {
            "Артикул продавца": "SKU-001",
            "Наименование": "Товар",
            "Материал": "Шёлк",     # already filled — must NOT be overwritten
            "Сезон": "",             # empty — should be filled
        }
    ]
    path = _make_wb_excel(tmp_path, rows)

    reader = WbExcelReader(path)
    products = reader.read_products()
    products[0]["ai_filled"] = {"Материал": "Хлопок", "Сезон": "Лето"}

    output = tmp_path / "output.xlsx"
    writer = WbExcelWriter(path)
    writer.write_filled(output, products)

    df_out = pd.read_excel(output, sheet_name=WB_PRODUCTS_SHEET, dtype=str)
    row = df_out.iloc[0]
    # Pre-filled value preserved
    assert row["Материал"] == "Шёлк"
    # Empty cell filled by AI
    assert row["Сезон"] == "Лето"


# ---------------------------------------------------------------------------
# 6. get_target_attributes returns only non-system columns
# ---------------------------------------------------------------------------

def test_get_target_attributes_returns_characteristic_columns(tmp_path):
    rows = [
        {
            "Артикул продавца": "SKU-001",
            "Наименование": "Товар",
            "Бренд": "Brand",
            "Материал": "Хлопок",
            "Страна производства": "Россия",
        }
    ]
    path = _make_wb_excel(tmp_path, rows)
    reader = WbExcelReader(path)
    products = reader.read_products()
    targets = reader.get_target_attributes(products)

    target_names = {t["name"] for t in targets}
    assert "Материал" in target_names
    assert "Страна производства" in target_names
    # System columns must not be in targets
    for sys_col in WB_SYSTEM_COLUMNS:
        assert sys_col not in target_names
    # Each target has required fields
    for t in targets:
        assert "id" in t
        assert "name" in t
        assert "type" in t
