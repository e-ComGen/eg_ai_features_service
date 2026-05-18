"""Tests for Ozon Excel parser/writer (app/services/excel/ozon_excel.py)."""
import io
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel.ozon_excel import (
    OZON_PRODUCTS_SHEET,
    OZON_VALIDATION_SHEET,
    OzonExcelReader,
    OzonExcelWriter,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _make_ozon_excel(
    tmp_path: Path,
    rows: list[dict],
    validation: dict[str, list[str]] | None = None,
) -> Path:
    """Create a minimal Ozon-style Excel file.

    Optionally adds a 'validation' sheet with allowed values.
    """
    path = tmp_path / "ozon_template.xlsx"
    wb = Workbook()

    # Products sheet
    ws_prod = wb.active
    ws_prod.title = OZON_PRODUCTS_SHEET
    if rows:
        headers = list(rows[0].keys())
        ws_prod.append(headers)
        for row in rows:
            ws_prod.append([row.get(h) for h in headers])

    # Validation sheet
    if validation is not None:
        ws_val = wb.create_sheet(title=OZON_VALIDATION_SHEET)
        col_names = list(validation.keys())
        ws_val.append(col_names)
        max_vals = max((len(v) for v in validation.values()), default=0)
        for i in range(max_vals):
            val_row = []
            for col in col_names:
                vals = validation[col]
                val_row.append(vals[i] if i < len(vals) else None)
            ws_val.append(val_row)

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# 1. Reader extracts validation sheet
# ---------------------------------------------------------------------------

def test_reader_extracts_validation_sheet(tmp_path):
    rows = [{"Артикул": "SKU-001", "Название товара": "Товар", "Цвет": ""}]
    validation = {
        "Цвет": ["Красный", "Синий", "Зелёный"],
        "Материал": ["Хлопок", "Полиэстер"],
    }
    path = _make_ozon_excel(tmp_path, rows, validation=validation)
    reader = OzonExcelReader(path)
    result = reader.read_validation()

    assert "Цвет" in result
    assert set(result["Цвет"]) == {"Красный", "Синий", "Зелёный"}
    assert "Материал" in result
    assert set(result["Материал"]) == {"Хлопок", "Полиэстер"}


# ---------------------------------------------------------------------------
# 2. Reader returns empty dict when no validation sheet
# ---------------------------------------------------------------------------

def test_reader_validation_when_no_validation_sheet_returns_empty(tmp_path):
    rows = [{"Артикул": "SKU-001", "Название товара": "Товар"}]
    path = _make_ozon_excel(tmp_path, rows, validation=None)
    reader = OzonExcelReader(path)
    result = reader.read_validation()

    assert result == {}


# ---------------------------------------------------------------------------
# 3. Reader extracts products
# ---------------------------------------------------------------------------

def test_reader_extracts_products(tmp_path):
    rows = [
        {
            "Артикул": "OZ-001",
            "Название товара": "Платье летнее",
            "Бренд": "SummerBrand",
            "Аннотация": "Лёгкое летнее платье",
            "Цвет": "",
            "Размер": "",
        },
        {
            "Артикул": "OZ-002",
            "Название товара": "Юбка",
            "Бренд": "AnotherBrand",
            "Аннотация": "Короткая юбка",
            "Цвет": "Синий",
            "Размер": "S",
        },
    ]
    path = _make_ozon_excel(tmp_path, rows)
    reader = OzonExcelReader(path)
    products = reader.read_products()

    assert len(products) == 2
    p0 = products[0]
    assert p0["sku"] == "OZ-001"
    assert p0["name"] == "Платье летнее"
    assert p0["brand"] == "SummerBrand"
    assert p0["description"] == "Лёгкое летнее платье"
    assert "row_index" in p0
    assert "_all_columns" in p0


# ---------------------------------------------------------------------------
# 4. Writer preserves validation sheet
# ---------------------------------------------------------------------------

def test_writer_preserves_validation_sheet(tmp_path):
    rows = [{"Артикул": "OZ-001", "Название товара": "Товар", "Цвет": ""}]
    validation = {"Цвет": ["Красный", "Синий"]}
    path = _make_ozon_excel(tmp_path, rows, validation=validation)

    reader = OzonExcelReader(path)
    products = reader.read_products()
    products[0]["ai_filled"] = {"Цвет": "Красный"}

    output = tmp_path / "output.xlsx"
    writer = OzonExcelWriter(path)
    writer.write_filled(output, products)

    assert output.exists()
    wb = load_workbook(output)
    # Both sheets preserved
    assert OZON_PRODUCTS_SHEET in wb.sheetnames
    assert OZON_VALIDATION_SHEET in wb.sheetnames
    # Value written into products sheet
    ws = wb[OZON_PRODUCTS_SHEET]
    headers = [cell.value for cell in ws[1]]
    цвет_idx = headers.index("Цвет") + 1  # 1-based
    assert ws.cell(row=2, column=цвет_idx).value == "Красный"


# ---------------------------------------------------------------------------
# 5. get_target_attributes uses validation for enum type + allowed_values
# ---------------------------------------------------------------------------

def test_get_target_attributes_uses_validation_for_enum(tmp_path):
    rows = [{"Артикул": "OZ-001", "Название товара": "Товар", "Цвет": "", "Вес (кг)": ""}]
    validation = {"Цвет": ["Красный", "Синий", "Зелёный"]}
    path = _make_ozon_excel(tmp_path, rows, validation=validation)

    reader = OzonExcelReader(path)
    products = reader.read_products()
    targets = reader.get_target_attributes(products)

    target_map = {t["name"]: t for t in targets}
    # Column with validation → enum type with allowed_values
    assert target_map["Цвет"]["type"] == "enum"
    assert set(target_map["Цвет"]["allowed_values"]) == {"Красный", "Синий", "Зелёный"}
    # Column without validation → text type
    assert target_map["Вес (кг)"]["type"] == "text"
    assert "allowed_values" not in target_map["Вес (кг)"] or target_map["Вес (кг)"].get("allowed_values") is None
