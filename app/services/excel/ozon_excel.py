"""Ozon Excel template parser/writer.

Имеет 2 вкладки:
- 'Шаблон для поставщика' — заполняемые данные
- 'validation' (скрытая) — словарь allowed values для каждой колонки

Лимит: 1000 товаров на файл, 25 МБ.
"""
import pandas as pd
from openpyxl import load_workbook
from pathlib import Path
from typing import List, Dict


OZON_PRODUCTS_SHEET = "Шаблон для поставщика"
OZON_VALIDATION_SHEET = "validation"

# Системные колонки Ozon (не characteristics)
OZON_SYSTEM_COLUMNS = {
    "Артикул", "Название товара", "Бренд", "Аннотация",
    "Изображения", "Ключевые слова", "Тип", "Категория",
    # TODO: уточнить полный список по мере встреч с реальными шаблонами
}


class OzonExcelReader:
    """Парсит Ozon Excel шаблон."""

    def __init__(self, file_path: "str | Path"):
        self.path = Path(file_path)

    def read_validation(self) -> "dict[str, list[str]]":
        """Возвращает {column_name: [allowed_values]} из скрытой вкладки validation.

        Это OZON-specific сокровище — словарь категорий в самом файле!
        """
        try:
            wb = load_workbook(self.path, read_only=True, data_only=True)
            if OZON_VALIDATION_SHEET not in wb.sheetnames:
                return {}
            ws = wb[OZON_VALIDATION_SHEET]
            # Структура validation sheet: row 0 = headers, дальше — значения per column
            validation: dict[str, list[str]] = {}
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return {}
            headers = rows[0]
            for col_idx, col_name in enumerate(headers):
                if not col_name:
                    continue
                values: set[str] = set()
                for row in rows[1:]:
                    if col_idx < len(row) and row[col_idx] is not None:
                        values.add(str(row[col_idx]).strip())
                validation[str(col_name).strip()] = sorted(values)
            return validation
        except Exception:
            return {}

    def read_products(self) -> "list[dict]":
        """Returns list of {row_index, sku, name, brand, description, characteristics, _all_columns}.

        Same pattern as WbExcelReader.read_products.
        """
        df = pd.read_excel(self.path, sheet_name=OZON_PRODUCTS_SHEET, dtype=str)
        df = df.dropna(subset=[df.columns[0]])

        # Load validation for allowed_values enrichment
        validation = self.read_validation()

        products = []
        for idx, row in df.iterrows():
            characteristics = {}
            for col in df.columns:
                if col not in OZON_SYSTEM_COLUMNS and pd.notna(row[col]):
                    characteristics[col] = row[col]

            products.append({
                "row_index": idx,
                "sku": row.get("Артикул"),
                "name": row.get("Название товара"),
                "brand": row.get("Бренд"),
                "description": row.get("Аннотация"),
                "characteristics": characteristics,
                "_all_columns": list(df.columns),
                "_validation": validation,  # allowed_values per column from validation sheet
            })
        return products

    def get_target_attributes(self, products: List[Dict]) -> List[Dict]:
        """Извлекает target characteristic columns с allowed_values из validation sheet."""
        if not products:
            return []
        all_cols = products[0]["_all_columns"]
        validation = products[0].get("_validation", {})
        targets = []
        target_id = 1
        for col in all_cols:
            if col in OZON_SYSTEM_COLUMNS:
                continue
            target: dict = {"id": target_id, "name": col, "type": "text"}
            # Если есть allowed_values из validation sheet — это enum
            if col in validation and validation[col]:
                target["type"] = "enum"
                target["allowed_values"] = validation[col]
            targets.append(target)
            target_id += 1
        return targets


class OzonExcelWriter:
    """Write-back с preservation структуры (включая validation sheet и стили)."""

    def __init__(self, original_path: "str | Path"):
        self.path = Path(original_path)

    def write_filled(
        self,
        output_path: "str | Path",
        products_with_filled: List[Dict],
    ) -> Path:
        """Загружает оригинальный файл openpyxl, заполняет AI-результатами, сохраняет.

        Использует openpyxl для preservation скрытых листов / стилей.
        """
        wb = load_workbook(self.path)
        ws = wb[OZON_PRODUCTS_SHEET]

        # Build header -> col index map (1-based)
        header_map: dict[str, int] = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value is not None:
                header_map[str(cell.value)] = col_idx

        for product in products_with_filled:
            ai_filled = product.get("ai_filled", {})
            # row_index is 0-based DataFrame index; +2 for 1-based + header row
            xl_row = product["row_index"] + 2
            for col_name, val in ai_filled.items():
                col_idx = header_map.get(col_name)
                if col_idx is None:
                    continue
                cell = ws.cell(row=xl_row, column=col_idx)
                if not cell.value:
                    cell.value = val

        wb.save(output_path)
        return Path(output_path)
